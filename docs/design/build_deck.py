#!/usr/bin/env python3
# 16:9 deck — readability pass (larger type, tighter layout). px on 1920x1080; pt = px/2
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
import math

W,H=1920,1080
T=dict(bg="FAFAF7",surf="FFFFFF",teal="2D6A6E",dteal="1E4A4D",apri="F2A878",ink="142433",
       sec="4A5B6B",mut="8597A4",bord="E5E7EB",hero="E6F0F0",lt="EAF1F0",ltteal="A9C7C7",
       paincard="F2F5F6",painink="3A4A55",area="D9E7E6",guide="D7DDE0",aibg="FFF0E6",aibd="C45A1A",band2="F3F6F6",cline="9AA7AD")
PS=["2D6A6E","356985","C0703F","6E5A86","4C8060"]; PT=["DCE9E8","DCE7EE","F6E5D7","E8E2F0","DCEBE0"]
slides=[]
def NS(): s=[("rect",0,0,W,H,T["bg"],None,0)]; slides.append(s); return s
def R(s,x,y,w,h,fill,rad=0,stroke=None): s.append(("rect",x,y,w,h,fill,stroke,rad))
def L(s,x1,y1,x2,y2,col,wpx): s.append(("line",x1,y1,x2,y2,col,wpx))
def O(s,x,y,w,h,fill,stroke): s.append(("oval",x,y,w,h,fill,stroke))
def AREA(s,pts,fill): s.append(("area",pts,fill))
def TX(s,x,y,w,h,runs,anchor='t',align='l',autofit=False):
    if isinstance(runs,str): runs=[(runs,16,False,T["ink"],0)]
    s.append(("text",x,y,w,h,runs,anchor,align,autofit))
def LINE2(s,x,y,w,h,lead,rest,size): s.append(("line2",x,y,w,h,lead,rest,size))
def RT(s,x,y,w,h,txt,size,color): s.append(("rtext",x,y,w,h,txt,size,color))
def DIA(s,x,y,w,h,fill,stroke,txt,size,color): s.append(("dia",x,y,w,h,fill,stroke,txt,size,color))
def CELL(s,x,y,w,h,fill,txt,size,color,bold=False,align='c',rad=12,stroke=None,pad=14):
    R(s,x,y,w,h,fill,rad,stroke); TX(s,x+pad,y,w-2*pad,h,[(txt,size,bold,color,0)],'m',align,True)

L0=100; R0=1820
# ===================== SLIDE 1 — OVERVIEW =====================
s=NS()
TX(s,L0,78,900,34,[("OVERVIEW",26,True,T["mut"],0)])
TX(s,L0,118,1720,100,[("Sona: Care Journey & Product Overview",80,True,T["teal"],0)])
TX(s,L0,236,1660,96,[("An AI practice partner that reduces admin burnout in UK private Speech and Language Therapy (SLT), giving clinicians more productive time for care across a client's first 30 days.",36,False,T["sec"],0)])
R(s,L0+2,366,160,6,T["apri"],3)
TX(s,L0,418,840,40,[("WHAT IS SONA?",32,True,T["teal"],0)])
TX(s,L0,478,830,420,[
  ("Sona is an AI practice partner for UK private Speech and Language Therapy (SLT).",34,False,T["ink"],0),
  ("It does the repetitive admin alongside the clinician: intake, consult prep, triage, the first session plan and family summary, all as drafts the clinician reviews.",34,False,T["ink"],18),
  ("Sona augments clinicians; it increases human productivity: less admin burnout, more time for care.",34,False,T["ink"],18)])
R(s,968,420,2,486,T["bord"],0)
TX(s,1016,418,820,40,[("HOW TO READ THIS",32,True,T["teal"],0)])
TX(s,1016,478,820,40,[("Three views, read left to right:",32,False,T["sec"],0)])
LINE2(s,1016,540,860,48,"Executive",": the journey and the value",40)
LINE2(s,1016,606,860,48,"Product",": the proposed solution",40)
LINE2(s,1016,672,860,48,"Process flow",": the To-Be swimlane",40)
TX(s,1016,762,820,36,[("WHO'S WHO",26,True,T["teal"],0)])
TX(s,1016,806,840,90,[("Client, parent or carer: the family or the adult client.\nClinician (SLT): the therapist. Practice: the private clinic.",28,False,T["sec"],0)])
R(s,L0,912,1720,140,T["lt"],16)
TX(s,L0+32,938,600,32,[("ANCHOR PRINCIPLES",24,True,T["teal"],0)])
for p,x in zip(["AI drafts; clinician decides","UK data residency + audit","Augment people, never replace them"],[140,713,1287]):
    TX(s,x,988,520,40,[(p,26,False,T["dteal"],0)])
for x in [673,1247]: R(s,x,982,1,52,T["bord"],0)

# ===================== SLIDE 2 — EXECUTIVE =====================
s=NS()
TX(s,L0,68,900,32,[("EXECUTIVE VIEW",24,True,T["mut"],0)])
TX(s,L0,104,1640,80,[("The care journey at a glance",72,True,T["teal"],0)])
TX(s,L0,210,1660,46,[("How a new SLT case feels today, and where Sona reduces admin burnout and adds value across five phases.",32,False,T["sec"],0)])
R(s,L0+2,272,150,6,T["apri"],3)
gap=22; cw=(R0-L0-4*gap)/5; step=cw+gap; xs=[L0+i*step for i in range(5)]; cx=[x+cw/2 for x in xs]
ph=["1 · Find & enquire","2 · Understand","3 · Meet & decide","4 · Plan & share","5 · Progress together"]
for i in range(5): CELL(s,xs[i],300,cw,84,PS[i],ph[i],30,"FFFFFF",True,'c',14)
TX(s,L0,408,1300,30,[("HOW EACH PHASE FEELS · one point per phase, family and clinician",24,True,T["mut"],0)])
lvl=[0.18,0.06,0.5,0.78,0.96]; base=560; py=[452+(1-l)*100 for l in lvl]
for i in range(5): L(s,cx[i],386,cx[i],py[i]-13,T["guide"],1.5)
AREA(s,[(cx[0],base)]+[(cx[i],py[i]) for i in range(5)]+[(cx[4],base)],T["area"])
L(s,L0,base,R0,base,T["bord"],1.5)
for i in range(4): L(s,cx[i],py[i],cx[i+1],py[i+1],T["teal"],5)
moods=["Anxious","Overwhelmed","Hopeful","Reassured","Confident"]
desc=["unsure where to start","so much to take in","a clear path forward","kept in the loop","confident and in control"]
for i in range(5):
    O(s,cx[i]-13,py[i]-13,26,26,PS[i],"FFFFFF")
    TX(s,xs[i],576,cw,32,[(moods[i],30,True,PS[i],0)],'m','c'); TX(s,xs[i],614,cw,28,[(desc[i],22,False,T["sec"],0)],'m','c')
TX(s,L0,664,900,30,[("TODAY'S PAIN · BY PHASE",24,True,T["mut"],0)])
pains=["Enquiries slip through the cracks","Chasing families to finish the form","Prep is rushed; the free consult under-delivers","Plans are hand-built; reports take ~3 weeks","Carryover is ad hoc (email or WhatsApp)"]
for i in range(5):
    R(s,xs[i],700,cw,118,T["paincard"],12,T["bord"]); O(s,xs[i]+18,718,11,11,PS[i],None)
    TX(s,xs[i]+18,700,cw-36,118,[(pains[i],26,False,T["painink"],0)],'m','c',True)
TX(s,L0,852,900,30,[("TARGET OUTCOMES",24,True,T["mut"],0)])
R(s,L0+2,888,100,6,T["apri"],3); TX(s,L0,902,860,70,[("≥ 4 hours a month",68,True,T["teal"],0)]); TX(s,L0,982,860,40,[("productive hours given back to each clinician",28,False,T["sec"],0)])
R(s,1012,888,100,6,T["apri"],3); TX(s,1010,902,820,70,[("Same-day",68,True,T["teal"],0)]); TX(s,1010,982,830,40,[("client summaries and reports, versus the ~3-week manual baseline",28,False,T["sec"],0)])

# ===================== SLIDE 3 — PRODUCT =====================
s=NS()
TX(s,L0,72,900,32,[("PRODUCT VIEW",24,True,T["mut"],0)])
TX(s,L0,108,1640,80,[("The proposed solution",72,True,T["teal"],0)])
TX(s,L0,214,1680,46,[("The first 30 days as nine stages grouped into five phases: what Sona does and the surface that delivers it.",30,False,T["sec"],0)])
lab=176; gx=L0+lab+14; ncol=9; cg=12; cwd=(R0-gx-(ncol-1)*cg)/ncol; st2=cwd+cg; colx=[gx+i*st2 for i in range(ncol)]; spn=[0,1,1,2,2,2,3,3,4]
stg=["1 · Referral","2 · Smart intake","3 · Intake review","4 · Consult prep","5 · Consultation","6 · Triage","7 · Session plan","8 · Family summary","9 · Carryover"]
wh=["Capture enquiry + audit trail","Adaptive intake + status + consent","One-screen client overview","AI prep brief (DRAFT)","Schedule consult (calendar, video)","Capture triage + rationale","AI session plan (DRAFT)","Tone-adjusted client summary","Resources + progress portal"]
sf=["Clinician web","Family or client app","Clinician web","Clinician web + AI","Clinician web + calendar","Clinician web","Clinician web + AI","Client app + AI","Client app + web"]
plab=["Find & enquire","Understand","Meet & decide","Plan & share","Progress together"]; spans={0:(0,0),1:(1,2),2:(3,5),3:(6,7),4:(8,8)}
yP,hP,yS,hS,yW,hW,yU,hU=296,56,360,66,434,210,652,118
TX(s,L0,yP,lab,hP,[("Phase",26,True,T["sec"],0)],'m','l'); TX(s,L0,yS,lab,hS,[("Stage",26,True,T["sec"],0)],'m','l')
TX(s,L0,yW,lab,hW,[("What Sona\ndoes",26,True,T["sec"],0)],'m','l'); TX(s,L0,yU,lab,hU,[("Surface /\nowner",26,True,T["sec"],0)],'m','l')
for k,(a,b) in spans.items():
    x0=colx[a]; x1=colx[b]+cwd; CELL(s,x0,yP,x1-x0,hP,PS[k],plab[k],26,"FFFFFF",True,'c',10)
for i in range(ncol):
    p=spn[i]; x=colx[i]
    CELL(s,x,yS,cwd,hS,PS[p],stg[i],22,"FFFFFF",True,'c',10,None,6)
    CELL(s,x,yW,cwd,hW,PT[p],wh[i],23,T["ink"],False,'c',10,None,12)
    CELL(s,x,yU,cwd,hU,PT[p],sf[i],22,T["ink"],False,'c',10,None,10)
TX(s,L0,800,900,30,[("WHERE SONA WINS",24,True,T["mut"],0)])
opp=["Own the first 30 days as one loop","Trust is the moat; the clinician signs every draft","Start paediatric, then generalise","UK data residency + clinical audit"]
ow=(R0-L0-3*20)/4; ostep=ow+20
for i,o in enumerate(opp): CELL(s,L0+i*ostep,836,ow,96,T["hero"],o,24,T["teal"],True,'c',12,None,18)
TX(s,L0,956,1720,30,[("Colour groups the nine stages into their five phases. · Generalised across SLTs (paediatric and adult).",22,False,T["mut"],0)])

# ===================== SLIDE 4 — PROCESS FLOW =====================
s=NS()
TX(s,L0,64,900,32,[("PROCESS FLOW",24,True,T["mut"],0)])
TX(s,L0,100,1500,80,[("How the journey runs with Sona",70,True,T["teal"],0)])
TX(s,L0,200,1720,40,[("A swimlane map of the proposed (To-Be) process: the basis for prototypes and screens.",30,False,T["sec"],0)])
top=276; lh=180; labW=170; bandX=100+labW; bandW=1820-bandX; ly=[top+i*lh+lh//2 for i in range(4)]
lanes=[("Client / Parent / Carer","C98A5E"),("Clinician / SLT","2D6A6E"),("Sona (app & AI)","356985"),("External (calendar and PMS)","4C8060")]
for i,(nm,col) in enumerate(lanes):
    y=top+i*lh; R(s,bandX,y,bandW,lh,(T["band2"] if i%2 else "FFFFFF"),0,T["bord"]); R(s,100,y,labW,lh,col,0); RT(s,100,y,labW,lh,nm,22,"FFFFFF")
NP={
 "start":(330,ly[0],176,82,"pill"),"end":(1740,ly[0],170,82,"pill"),
 "comp":(516,ly[0],172,76,"box"),"recv":(1556,ly[0],172,76,"box"),
 "dec":(1020,ly[1],188,108,"dia"),"rev":(672,ly[1],172,76,"box"),"cons":(854,ly[1],172,76,"box"),
 "strat":(1240,ly[1],172,76,"box"),"edit":(1452,ly[1],172,76,"box"),"send":(1644,ly[1],172,76,"box"),
 "cap":(330,ly[2],176,76,"box"),"prep":(540,ly[2],172,76,"box"),"plan":(1196,ly[2],172,76,"box"),
 "sum":(1452,ly[2],172,76,"box"),"carry":(1670,ly[2],172,76,"box"),
 "cal":(770,ly[3],172,76,"box"),"pms":(1670,ly[3],172,76,"box")}
LB={"start":"Enquiry arrives","end":"Ongoing care","comp":"Complete intake & consent","recv":"Receive summary (app)",
 "dec":"Ongoing therapy?","rev":"Review overview + prep","cons":"Consultation (~20 min)","strat":"Strategy advice & home tips",
 "edit":"Edit & sign plan","send":"Review & send summary","cap":"Capture enquiry + smart intake","prep":"AI prep brief (DRAFT)",
 "plan":"AI session plan (DRAFT)","sum":"AI family summary (DRAFT)","carry":"Carryover + progress portal","cal":"Calendar and video link","pms":"PMS sync (vision)"}
for n,(cxN,cyN,w,h,kind) in NP.items():
    x=cxN-w/2; y=cyN-h/2
    if kind=="box": CELL(s,x,y,w,h,T["hero"],LB[n],22,T["ink"],False,'c',12,T["teal"],10)
    elif kind=="pill": CELL(s,x,y,w,h,T["surf"],LB[n],22,T["ink"],True,'c',h/2,"C98A5E",10)
    elif kind=="dia": DIA(s,x,y,w,h,T["aibg"],T["aibd"],LB[n],22,T["ink"])
def CONN(a,b,label=""):
    ax,ay,aw,ah,_=NP[a]; bx,by,bw,bh,_=NP[b]; dx=bx-ax; dy=by-ay
    if abs(dx)>=abs(dy): x1=ax+(aw/2 if dx>0 else -aw/2); y1=ay; x2=bx+(-bw/2 if dx>0 else bw/2); y2=by; orient='h'
    else: x1=ax; y1=ay+(ah/2 if dy>0 else -ah/2); x2=bx; y2=by+(-bh/2 if dy>0 else bh/2); orient='v'
    s.append(("conn",x1,y1,x2,y2,orient,label))
for a,b,lbl in [("start","cap",""),("cap","comp",""),("comp","prep",""),("prep","rev",""),("cal","cons",""),("rev","cons",""),("cons","dec",""),("dec","plan","Yes"),("dec","strat","No"),("plan","edit",""),("edit","sum",""),("strat","sum",""),("sum","send",""),("send","recv",""),("send","carry",""),("recv","end",""),("carry","pms","")]:
    CONN(a,b,lbl)
TX(s,bandX,1016,1720,30,[("Shapes: rounded = task · diamond = decision · pill = start or end. DRAFT = AI artefact the clinician reviews and signs.",22,False,T["mut"],0)])

# ===================== RENDER PPTX =====================
def C(h): return RGBColor.from_string(h)
def IN(px): return Inches(px/144.0)
def PTf(px): return Pt(px/2.0)
AL={'l':PP_ALIGN.LEFT,'c':PP_ALIGN.CENTER,'r':PP_ALIGN.RIGHT}; AN={'t':MSO_ANCHOR.TOP,'m':MSO_ANCHOR.MIDDLE}
def fl(shp,fill,stroke,sw=1.0):
    if fill is None: shp.fill.background()
    else: shp.fill.solid(); shp.fill.fore_color.rgb=C(fill)
    if stroke is None: shp.line.fill.background()
    else: shp.line.color.rgb=C(stroke); shp.line.width=Pt(sw)
    shp.shadow.inherit=False
def settext(shp,txt,size,color,bold,align='c'):
    tf=shp.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.MIDDLE; tf.auto_size=MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    for m in ('margin_left','margin_right','margin_top','margin_bottom'): setattr(tf,m,Pt(2))
    p=tf.paragraphs[0]; p.alignment=AL[align]; r=p.add_run(); r.text=txt; r.font.size=PTf(size); r.font.bold=bold; r.font.color.rgb=C(color); r.font.name="Inter"
prs=Presentation(); prs.slide_width=IN(W); prs.slide_height=IN(H); BL=prs.slide_layouts[6]
for ops in slides:
    sl=prs.slides.add_slide(BL)
    for op in ops:
        k=op[0]
        if k=="rect":
            _,x,y,w,h,fill,stroke,rad=op; shp=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rad else MSO_SHAPE.RECTANGLE,IN(x),IN(y),IN(w),IN(h)); fl(shp,fill,stroke)
            if rad:
                try: shp.adjustments[0]=max(0,min(0.5,rad/min(w,h)))
                except Exception: pass
        elif k=="oval":
            _,x,y,w,h,fill,stroke=op; shp=sl.shapes.add_shape(MSO_SHAPE.OVAL,IN(x),IN(y),IN(w),IN(h)); fl(shp,fill,stroke,1.5)
        elif k=="dia":
            _,x,y,w,h,fill,stroke,txt,size,color=op; shp=sl.shapes.add_shape(MSO_SHAPE.DIAMOND,IN(x),IN(y),IN(w),IN(h)); fl(shp,fill,stroke,1.5); settext(shp,txt,size,color,True)
        elif k=="area":
            _,pts,fill=op
            try:
                e=[(int(IN(px)),int(IN(py))) for (px,py) in pts]
                fb=sl.shapes.build_freeform(e[0][0],e[0][1],scale=1.0); fb.add_line_segments(e[1:],close=True); shp=fb.convert_to_shape(); fl(shp,fill,None)
            except Exception: pass
        elif k=="line":
            _,x1,y1,x2,y2,col,wpx=op; cn=sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,IN(x1),IN(y1),IN(x2),IN(y2)); cn.line.color.rgb=C(col); cn.line.width=Pt(wpx/2.0); cn.shadow.inherit=False
        elif k=="conn":
            _,x1,y1,x2,y2,orient,label=op
            cn=sl.shapes.add_connector(MSO_CONNECTOR.ELBOW,IN(x1),IN(y1),IN(x2),IN(y2)); cn.line.color.rgb=C(T["cline"]); cn.line.width=Pt(1.5); cn.shadow.inherit=False
            ln=cn.line._get_or_add_ln(); ln.append(ln.makeelement(qn('a:tailEnd'),{'type':'triangle','w':'med','len':'med'}))
            if label:
                mx=(x1+x2)/2; my=(y1+y2)/2; tb=sl.shapes.add_textbox(IN(mx-32),IN(my-24),IN(64),IN(26)); pp=tb.text_frame.paragraphs[0]; pp.alignment=PP_ALIGN.CENTER; rr=pp.add_run(); rr.text=label; rr.font.size=PTf(22); rr.font.bold=True; rr.font.color.rgb=C(T["sec"]); rr.font.name="Inter"
        elif k=="rtext":
            _,x,y,w,h,txt,size,color=op; tb=sl.shapes.add_textbox(IN(x),IN(y),IN(w),IN(h)); tb.rotation=270; tf=tb.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.MIDDLE
            p=tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER; r=p.add_run(); r.text=txt; r.font.size=PTf(size); r.font.bold=True; r.font.color.rgb=C(color); r.font.name="Inter"
        elif k=="line2":
            _,x,y,w,h,lead,rest,size=op; tb=sl.shapes.add_textbox(IN(x),IN(y),IN(w),IN(h)); tf=tb.text_frame; tf.word_wrap=True; p=tf.paragraphs[0]; p.alignment=PP_ALIGN.LEFT
            r1=p.add_run(); r1.text=lead; r1.font.size=PTf(size); r1.font.bold=True; r1.font.color.rgb=C(T["teal"]); r1.font.name="Inter"
            r2=p.add_run(); r2.text=rest; r2.font.size=PTf(size); r2.font.bold=False; r2.font.color.rgb=C(T["ink"]); r2.font.name="Inter"
        elif k=="text":
            _,x,y,w,h,runs,anchor,align,autofit=op; tb=sl.shapes.add_textbox(IN(x),IN(y),IN(w),IN(h)); tf=tb.text_frame; tf.word_wrap=True; tf.vertical_anchor=AN[anchor]
            if autofit: tf.auto_size=MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
            for m in ('margin_left','margin_right','margin_top','margin_bottom'): setattr(tf,m,Pt(2))
            for i,(txt,size,bold,color,spb) in enumerate(runs):
                p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=AL[align]
                if spb: p.space_before=PTf(spb)
                p.space_after=Pt(0); r=p.add_run(); r.text=txt; r.font.size=PTf(size); r.font.bold=bold; r.font.color.rgb=C(color); r.font.name="Inter"
    # footer per slide
    fy=1046
    tb=sl.shapes.add_textbox(IN(100),IN(fy),IN(900),IN(26)); rp=tb.text_frame.paragraphs[0]; rr=rp.add_run(); rr.text="Sona: Care Journey & Product Overview"; rr.font.size=PTf(20); rr.font.color.rgb=C(T["mut"]); rr.font.name="Inter"
    tb2=sl.shapes.add_textbox(IN(1620),IN(fy),IN(200),IN(26)); rp2=tb2.text_frame.paragraphs[0]; rp2.alignment=PP_ALIGN.RIGHT; rr2=rp2.add_run(); rr2.text=f"{slides.index(ops)+1} / 4"; rr2.font.size=PTf(20); rr2.font.color.rgb=C(T["mut"]); rr2.font.name="Inter"
prs.save("/tmp/Sona_Care_Journey_Deck.pptx"); print("pptx saved",len(slides))

# ===================== RENDER PNG PREVIEW =====================
FR="/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"; FB="/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"
def rgb(h): return tuple(int(h[i:i+2],16) for i in (0,2,4))
def font(px,bold): return ImageFont.truetype(FB if bold else FR, max(6,int(round(px))))
def wrap(d,txt,f,mw):
    out=[]
    for seg in txt.split("\n"):
        words=seg.split(); cur=""
        if not words: out.append(""); continue
        for wd in words:
            t=(cur+" "+wd).strip()
            if d.textlength(t,font=f)<=mw or not cur: cur=t
            else: out.append(cur); cur=wd
        out.append(cur)
    return out
def layout(d,runs,mw,sc):
    blks=[]
    for (txt,size,bold,color,spb) in runs:
        f=font(size*sc,bold); ls=wrap(d,txt,f,mw); a,de=f.getmetrics(); lh=a+de+int(4*sc); blks.append((ls,f,rgb(color),int(spb*sc),lh))
    return blks,sum(spb+lh*len(ls) for ls,f,c,spb,lh in blks)
def drawtext(d,x,y,w,h,runs,anchor,align,autofit):
    mw=w-4; sc=1.0; blks,tot=layout(d,runs,mw,sc)
    if autofit and tot>h-2 and tot>0: sc=max(0.4,(h-2)/tot); blks,tot=layout(d,runs,mw,sc)
    cy=y+1 if anchor=='t' else y+max(0,(h-tot)/2)
    for ls,f,col,spb,lh in blks:
        cy+=spb
        for ln in ls:
            tw=d.textlength(ln,font=f); tx=x+(w-tw)/2 if align=='c' else (x+w-2-tw if align=='r' else x+2)
            d.text((tx,cy),ln,font=f,fill=col); cy+=lh
def arrow(d,p,ang,col,sz=12):
    x,y=p; a=math.radians(ang); d.polygon([(x,y),(x-sz*math.cos(a-0.5),y-sz*math.sin(a-0.5)),(x-sz*math.cos(a+0.5),y-sz*math.sin(a+0.5))],fill=col)
imgs=[]
for idx,ops in enumerate(slides):
    im=Image.new("RGB",(W,H),(255,255,255)); d=ImageDraw.Draw(im)
    for op in ops:
        k=op[0]
        if k=="rect":
            _,x,y,w,h,fill,stroke,rad=op; box=[x,y,x+w,y+h]
            if rad: d.rounded_rectangle(box,radius=min(rad,h/2),fill=rgb(fill) if fill else None,outline=rgb(stroke) if stroke else None,width=2 if stroke else 1)
            else: d.rectangle(box,fill=rgb(fill) if fill else None,outline=rgb(stroke) if stroke else None,width=2 if stroke else 1)
        elif k=="oval":
            _,x,y,w,h,fill,stroke=op; d.ellipse([x,y,x+w,y+h],fill=rgb(fill) if fill else None,outline=rgb(stroke) if stroke else None,width=3 if stroke else 1)
        elif k=="area":
            _,pts,fill=op; d.polygon([(px,py) for px,py in pts],fill=rgb(fill))
        elif k=="dia":
            _,x,y,w,h,fill,stroke,txt,size,color=op; d.polygon([(x+w/2,y),(x+w,y+h/2),(x+w/2,y+h),(x,y+h/2)],fill=rgb(fill),outline=rgb(stroke)); drawtext(d,x+16,y,w-32,h,[(txt,size,True,color,0)],'m','c',True)
        elif k=="line":
            _,x1,y1,x2,y2,col,wpx=op; d.line([x1,y1,x2,y2],fill=rgb(col),width=max(1,int(wpx)))
        elif k=="conn":
            _,x1,y1,x2,y2,orient,label=op; col=rgb(T["cline"])
            if orient=='h': mx=(x1+x2)/2; pts=[(x1,y1),(mx,y1),(mx,y2),(x2,y2)]; ang=0 if x2>=mx else 180
            else: my=(y1+y2)/2; pts=[(x1,y1),(x1,my),(x2,my),(x2,y2)]; ang=90 if y2>=my else 270
            d.line(pts,fill=col,width=2,joint="curve"); arrow(d,(x2,y2),ang,col)
            if label: drawtext(d,(x1+x2)/2-34,(y1+y2)/2-24,68,26,[(label,22,True,T["sec"],0)],'m','c',False)
        elif k=="rtext":
            _,x,y,w,h,txt,size,color=op; th=int(h); tw=int(170); ti=Image.new("RGBA",(th,tw),(0,0,0,0)); td=ImageDraw.Draw(ti)
            f=font(size,True); lines=wrap(td,txt,f,th-10); a,de=f.getmetrics(); llh=a+de+2; cy=(tw-llh*len(lines))/2
            for ln in lines: lw=td.textlength(ln,font=f); td.text(((th-lw)/2,cy),ln,font=f,fill=rgb(color)); cy+=llh
            ti=ti.rotate(90,expand=True); im.paste(ti,(int(x),int(y)),ti)
        elif k=="line2":
            _,x,y,w,h,lead,rest,size=op; fL=font(size,True); fR=font(size,False); a,de=fL.getmetrics(); yy=y+max(0,(h-(a+de))/2)
            d.text((x,yy),lead,font=fL,fill=rgb(T["teal"])); lw=d.textlength(lead,font=fL); d.text((x+lw,yy),rest,font=fR,fill=rgb(T["ink"]))
        elif k=="text":
            _,x,y,w,h,runs,anchor,align,autofit=op; drawtext(d,x,y,w,h,runs,anchor,align,autofit)
    drawtext(d,100,1046,900,26,[("Sona: Care Journey & Product Overview",20,False,T["mut"],0)],'t','l',False)
    drawtext(d,1620,1046,200,26,[(f"{idx+1} / 4",20,False,T["mut"],0)],'t','r',False)
    imgs.append(im)
for i,im in enumerate(imgs): im.save(f"/tmp/n_slide_{i+1}.png")
gap=28; sheet=Image.new("RGB",(W,H*len(imgs)+gap*(len(imgs)+1)),(220,220,220)); yy=gap
for im in imgs: sheet.paste(im,(0,yy)); yy+=H+gap
sheet.save("/tmp/n_deck_preview.png")
import img2pdf
with open("/tmp/Sona_Care_Journey_Deck.pdf","wb") as f: f.write(img2pdf.convert([f"/tmp/n_slide_{i}.png" for i in range(1,5)]))
print("preview + pdf saved")
